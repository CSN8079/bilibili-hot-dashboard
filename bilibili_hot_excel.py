#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""定时抓取哔哩哔哩全站热门排行并写入 Excel。"""

from __future__ import annotations

import argparse
import base64
import configparser
import dataclasses
import html
import json
import logging
import os
import re
import shutil
import subprocess
import sys
import threading
import time
from dataclasses import dataclass
from datetime import datetime, timedelta
from pathlib import Path
from typing import Any, Sequence
from urllib.error import HTTPError, URLError
from urllib.parse import parse_qsl, urlencode, urlsplit, urlunsplit
from urllib.request import Request, urlopen

from dashboard_server import DashboardHTTPServer, DashboardRuntime, start_dashboard_server
from snapshot_store import (
    PendingQueueError,
    append_pending_snapshot,
    build_snapshot,
    clear_pending_snapshots,
    read_pending_snapshots,
    snapshot_rows,
)

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
except ImportError:
    print(
        "缺少依赖 openpyxl。请先运行：python -m pip install -r requirements.txt",
        file=sys.stderr,
    )
    raise SystemExit(2)


BASE_DIR = Path(__file__).resolve().parent
DEFAULT_CONFIG = BASE_DIR / "config.ini"
DEFAULT_OUTPUT = BASE_DIR / "bilibili_hot_videos.xlsx"
DEFAULT_PENDING_FILE = BASE_DIR / "pending_updates.jsonl"
DEFAULT_DASHBOARD_HTML = BASE_DIR / "bilibili_hot_dashboard.html"
DEFAULT_DASHBOARD_DATA = BASE_DIR / "dashboard_data.json"
DEFAULT_DASHBOARD_SHORTCUT = BASE_DIR / "bilibili_hot_dashboard.url"
DEFAULT_API_URL = "https://api.bilibili.com/x/web-interface/ranking/v2"
DEFAULT_USER_AGENT = (
    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
    "AppleWebKit/537.36 (KHTML, like Gecko) "
    "Chrome/124.0.0.0 Safari/537.36"
)

OLD_HEADERS_V1 = (
    "抓取时间",
    "排名",
    "视频名称",
    "分区",
    "播放量",
    "点赞量",
    "投币量",
    "收藏量",
    "BV号",
    "UP主",
    "视频链接",
)
OLD_HEADERS_V2 = (
    "抓取时间",
    "排名",
    "封面链接",
    "视频名称",
    "分区",
    "播放量",
    "点赞量",
    "投币量",
    "收藏量",
    "转发量",
    "BV号",
    "UP主",
    "视频链接",
)
HEADERS = (
    "抓取时间",
    "排名",
    "排名变化",
    "封面链接",
    "视频名称",
    "分区",
    "播放量",
    "点赞量",
    "投币量",
    "收藏量",
    "转发量",
    "BV号",
    "UP主",
    "视频链接",
)
COLUMN_WIDTHS = (20, 8, 12, 42, 54, 15, 14, 12, 12, 12, 12, 16, 20, 44)
NUMERIC_COLUMNS = (7, 8, 9, 10, 11)
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(color="FFFFFF", bold=True)
HEADER_ALIGNMENT = Alignment(horizontal="center", vertical="center")
THIN_BORDER = Border(
    left=Side(style="thin", color="D9E1F2"),
    right=Side(style="thin", color="D9E1F2"),
    top=Side(style="thin", color="D9E1F2"),
    bottom=Side(style="thin", color="D9E1F2"),
)


class ConfigError(ValueError):
    """配置项不合法。"""


class BilibiliApiError(RuntimeError):
    """哔哩哔哩接口返回业务错误。"""


class ExcelWriteError(RuntimeError):
    """Excel 写入失败，locked 表示文件很可能正被 Excel 占用。"""

    def __init__(self, message: str, locked: bool = False) -> None:
        super().__init__(message)
        self.locked = locked


_EXCEL_NOTIFICATION_LAST_SENT = 0.0


@dataclass(frozen=True)
class Settings:
    interval_minutes: float
    output_file: Path
    pending_file: Path
    top_n: int
    api_url: str
    ranking_type: str
    ranking_rid: str
    request_timeout: float
    retry_times: int
    retry_delay: float
    user_agent: str
    referer: str
    cookie: str
    log_level: str
    dashboard_enabled: bool
    dashboard_host: str
    dashboard_port: int
    dashboard_refresh_seconds: int
    dashboard_html: Path
    dashboard_data: Path
    dashboard_shortcut: Path
    notification_enabled: bool
    notification_cooldown_seconds: int


def parse_args(argv: Sequence[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="每隔指定时间抓取哔哩哔哩全站热门排行，并写入 Excel。",
    )
    parser.add_argument(
        "--config",
        type=Path,
        help="配置文件路径，默认使用程序目录中的 config.ini。",
    )
    parser.add_argument(
        "--interval-minutes",
        type=float,
        help="抓取间隔（分钟），覆盖配置文件。",
    )
    parser.add_argument("--output", type=Path, help="Excel 输出路径，覆盖配置文件。")
    parser.add_argument("--top-n", type=int, help="每次保存的视频数量（1-100）。")
    parser.add_argument("--timeout", type=float, help="单次网络请求超时秒数。")
    parser.add_argument("--retry-times", type=int, help="失败重试次数。")
    parser.add_argument("--retry-delay", type=float, help="首次重试等待秒数。")
    parser.add_argument(
        "--cookie",
        help="可选的 B 站 Cookie；通常不需要。建议使用配置文件或环境变量。",
    )
    parser.add_argument(
        "--log-level",
        choices=("DEBUG", "INFO", "WARNING", "ERROR"),
        help="日志级别，覆盖配置文件。",
    )
    parser.add_argument("--pending-file", type=Path, help="待补写数据队列文件。")
    parser.add_argument(
        "--no-notifications",
        action="store_true",
        help="Excel 被占用时不弹出 Windows 系统提示。",
    )
    parser.add_argument(
        "--notification-cooldown",
        type=int,
        help="相同通知的最短间隔秒数，0 表示每次都提示。",
    )
    parser.add_argument(
        "--no-dashboard",
        action="store_true",
        help="不生成 HTML 看板，也不启动本地看板服务。",
    )
    parser.add_argument("--dashboard-host", help="看板监听地址，默认 127.0.0.1。")
    parser.add_argument("--dashboard-port", type=int, help="看板监听端口，默认 8765。")
    parser.add_argument("--dashboard-refresh", type=int, help="网页自动刷新间隔秒数。")
    parser.add_argument(
        "--flush-pending",
        action="store_true",
        help="不抓取新数据，仅尝试把待补写快照写入 Excel。",
    )
    parser.add_argument(
        "--once",
        action="store_true",
        help="只抓取一次后退出，适用于 Windows 任务计划程序。",
    )
    parser.add_argument(
        "--print-config",
        action="store_true",
        help="打印最终生效的配置后退出，Cookie 会脱敏。",
    )
    return parser.parse_args(argv)


def _read_config(config_path: Path) -> configparser.ConfigParser:
    config = configparser.ConfigParser(interpolation=None)
    config.read(config_path, encoding="utf-8-sig")
    return config


def _config_value(
    config: configparser.ConfigParser,
    section: str,
    option: str,
    fallback: str,
) -> str:
    if config.has_section(section) and config.has_option(section, option):
        return config.get(section, option, fallback=fallback).strip()
    return fallback


def _to_float(value: str, name: str) -> float:
    try:
        return float(value)
    except (TypeError, ValueError) as exc:
        raise ConfigError(f"{name} 必须是数字，当前值为：{value!r}") from exc


def _to_int(value: str, name: str) -> int:
    try:
        return int(value)
    except (TypeError, ValueError) as exc:
        raise ConfigError(f"{name} 必须是整数，当前值为：{value!r}") from exc


def _resolve_path(value: str, base_dir: Path) -> Path:
    path = Path(value).expanduser()
    if not path.is_absolute():
        path = base_dir / path
    return path.resolve()


def _to_bool(value: str, name: str) -> bool:
    normalized = str(value).strip().lower()
    if normalized in {"1", "true", "yes", "on"}:
        return True
    if normalized in {"0", "false", "no", "off"}:
        return False
    raise ConfigError(f"{name} 必须是 true/false、yes/no 或 1/0。")


def load_settings(args: argparse.Namespace) -> tuple[Settings, Path]:
    if args.config:
        config_path = args.config.expanduser()
        if not config_path.is_absolute():
            config_path = (Path.cwd() / config_path).resolve()
    else:
        config_path = DEFAULT_CONFIG

    config = _read_config(config_path)
    config_dir = config_path.parent if config_path.exists() else BASE_DIR

    interval_raw = (
        str(args.interval_minutes)
        if args.interval_minutes is not None
        else _config_value(config, "schedule", "interval_minutes", "10")
    )
    output_raw = (
        str(args.output)
        if args.output is not None
        else _config_value(config, "excel", "output_file", str(DEFAULT_OUTPUT))
    )
    pending_raw = (
        str(args.pending_file)
        if args.pending_file is not None
        else _config_value(
            config,
            "storage",
            "pending_file",
            str(DEFAULT_PENDING_FILE),
        )
    )
    top_n_raw = (
        str(args.top_n)
        if args.top_n is not None
        else _config_value(config, "excel", "top_n", "100")
    )
    timeout_raw = (
        str(args.timeout)
        if args.timeout is not None
        else _config_value(config, "api", "request_timeout", "20")
    )
    retry_times_raw = (
        str(args.retry_times)
        if args.retry_times is not None
        else _config_value(config, "api", "retry_times", "3")
    )
    retry_delay_raw = (
        str(args.retry_delay)
        if args.retry_delay is not None
        else _config_value(config, "api", "retry_delay", "5")
    )
    dashboard_host = (
        args.dashboard_host
        or _config_value(config, "dashboard", "host", "127.0.0.1")
    )
    dashboard_port_raw = (
        str(args.dashboard_port)
        if args.dashboard_port is not None
        else _config_value(config, "dashboard", "port", "8765")
    )
    dashboard_refresh_raw = (
        str(args.dashboard_refresh)
        if args.dashboard_refresh is not None
        else _config_value(config, "dashboard", "refresh_seconds", "10")
    )
    dashboard_html_raw = _config_value(
        config,
        "dashboard",
        "html_file",
        str(DEFAULT_DASHBOARD_HTML),
    )
    dashboard_data_raw = _config_value(
        config,
        "dashboard",
        "data_file",
        str(DEFAULT_DASHBOARD_DATA),
    )
    dashboard_shortcut_raw = _config_value(
        config,
        "dashboard",
        "shortcut_file",
        str(DEFAULT_DASHBOARD_SHORTCUT),
    )
    notification_cooldown_raw = (
        str(args.notification_cooldown)
        if args.notification_cooldown is not None
        else _config_value(config, "notifications", "cooldown_seconds", "600")
    )
    dashboard_enabled = not args.no_dashboard and _to_bool(
        _config_value(config, "dashboard", "enabled", "true"),
        "dashboard.enabled",
    )

    interval_minutes = _to_float(interval_raw, "interval_minutes")
    top_n = _to_int(top_n_raw, "top_n")
    request_timeout = _to_float(timeout_raw, "request_timeout")
    retry_times = _to_int(retry_times_raw, "retry_times")
    retry_delay = _to_float(retry_delay_raw, "retry_delay")
    dashboard_port = _to_int(dashboard_port_raw, "dashboard.port")
    dashboard_refresh_seconds = _to_int(
        dashboard_refresh_raw,
        "dashboard.refresh_seconds",
    )
    notification_cooldown_seconds = _to_int(
        notification_cooldown_raw,
        "notifications.cooldown_seconds",
    )
    notification_enabled = not args.no_notifications and _to_bool(
        _config_value(config, "notifications", "enabled", "true"),
        "notifications.enabled",
    )

    if interval_minutes <= 0:
        raise ConfigError("interval_minutes 必须大于 0。")
    if not 1 <= top_n <= 100:
        raise ConfigError("top_n 必须在 1 到 100 之间。")
    if request_timeout <= 0:
        raise ConfigError("request_timeout 必须大于 0。")
    if retry_times < 0:
        raise ConfigError("retry_times 不能小于 0。")
    if retry_delay < 0:
        raise ConfigError("retry_delay 不能小于 0。")
    if not 1 <= dashboard_port <= 65535:
        raise ConfigError("dashboard.port 必须在 1 到 65535 之间。")
    if not 2 <= dashboard_refresh_seconds <= 3600:
        raise ConfigError("dashboard.refresh_seconds 必须在 2 到 3600 之间。")
    if not dashboard_host:
        raise ConfigError("dashboard.host 不能为空。")
    if notification_cooldown_seconds < 0:
        raise ConfigError("notifications.cooldown_seconds 不能小于 0。")

    cookie = (
        args.cookie
        if args.cookie is not None
        else os.environ.get(
            "BILIBILI_COOKIE",
            _config_value(config, "optional_auth", "cookie", ""),
        )
    )

    settings = Settings(
        interval_minutes=interval_minutes,
        output_file=_resolve_path(output_raw, config_dir),
        pending_file=_resolve_path(pending_raw, config_dir),
        top_n=top_n,
        api_url=_config_value(config, "api", "base_url", DEFAULT_API_URL),
        ranking_type=_config_value(config, "api", "type", "all"),
        ranking_rid=_config_value(config, "api", "rid", "0"),
        request_timeout=request_timeout,
        retry_times=retry_times,
        retry_delay=retry_delay,
        user_agent=_config_value(config, "api", "user_agent", DEFAULT_USER_AGENT),
        referer=_config_value(
            config,
            "api",
            "referer",
            "https://www.bilibili.com/v/popular/rank/all",
        ),
        cookie=cookie.strip(),
        log_level=(
            args.log_level
            or _config_value(config, "logging", "level", "INFO").upper()
        ),
        dashboard_enabled=dashboard_enabled,
        dashboard_host=dashboard_host,
        dashboard_port=dashboard_port,
        dashboard_refresh_seconds=dashboard_refresh_seconds,
        dashboard_html=_resolve_path(dashboard_html_raw, config_dir),
        dashboard_data=_resolve_path(dashboard_data_raw, config_dir),
        dashboard_shortcut=_resolve_path(dashboard_shortcut_raw, config_dir),
        notification_enabled=notification_enabled,
        notification_cooldown_seconds=notification_cooldown_seconds,
    )

    if not settings.api_url:
        raise ConfigError("api.base_url 不能为空。")
    if settings.log_level not in {"DEBUG", "INFO", "WARNING", "ERROR"}:
        raise ConfigError("logging.level 必须是 DEBUG、INFO、WARNING 或 ERROR。")

    return settings, config_path


def public_settings(settings: Settings) -> dict[str, Any]:
    result = dataclasses.asdict(settings)
    for key in (
        "output_file",
        "pending_file",
        "dashboard_html",
        "dashboard_data",
        "dashboard_shortcut",
    ):
        result[key] = str(getattr(settings, key))
    result["cookie"] = "<已设置>" if settings.cookie else "<未设置>"
    return result


def configure_logging(level: str) -> None:
    logging.basicConfig(
        level=getattr(logging, level),
        format="%(asctime)s | %(levelname)s | %(message)s",
        datefmt="%Y-%m-%d %H:%M:%S",
    )


def build_api_url(settings: Settings) -> str:
    parts = urlsplit(settings.api_url)
    query = dict(parse_qsl(parts.query, keep_blank_values=True))
    query["rid"] = settings.ranking_rid
    query["type"] = settings.ranking_type
    return urlunsplit(
        (parts.scheme, parts.netloc, parts.path, urlencode(query), parts.fragment)
    )


def fetch_ranking(settings: Settings) -> dict[str, Any]:
    url = build_api_url(settings)
    headers = {
        "Accept": "application/json, text/plain, */*",
        "Accept-Language": "zh-CN,zh;q=0.9,en;q=0.7",
        "Referer": settings.referer,
        "User-Agent": settings.user_agent,
    }
    if settings.cookie:
        headers["Cookie"] = settings.cookie

    last_error: Exception | None = None
    for attempt in range(settings.retry_times + 1):
        try:
            request = Request(url, headers=headers, method="GET")
            with urlopen(request, timeout=settings.request_timeout) as response:
                raw = response.read()
            payload = json.loads(raw.decode("utf-8"))
            code = payload.get("code")
            if code != 0:
                message = payload.get("message") or payload.get("msg") or "未知错误"
                raise BilibiliApiError(f"B站接口返回 code={code}，message={message}")
            return payload
        except (HTTPError, URLError, TimeoutError, OSError, json.JSONDecodeError, BilibiliApiError) as exc:
            last_error = exc
            if attempt >= settings.retry_times:
                break
            wait_seconds = settings.retry_delay * (2**attempt)
            logging.warning(
                "抓取失败（第 %d/%d 次）：%s；%.1f 秒后重试。",
                attempt + 1,
                settings.retry_times + 1,
                exc,
                wait_seconds,
            )
            time.sleep(wait_seconds)

    raise RuntimeError(f"连续抓取失败：{last_error}") from last_error


def clean_text(value: Any) -> str:
    text = html.unescape(str(value or ""))
    text = re.sub(r"<[^>]+>", "", text)
    text = re.sub(r"[\x00-\x08\x0B\x0C\x0E-\x1F]", "", text)
    return text.strip()


def safe_excel_text(value: Any) -> str:
    text = clean_text(value)
    if text.startswith(("=", "+", "-", "@")):
        return "'" + text
    return text


def normalize_cover_url(value: Any) -> str:
    url = clean_text(value)
    if url.startswith("//"):
        return "https:" + url
    if url.startswith("http://"):
        return "https://" + url[len("http://") :]
    return url


def to_int(value: Any) -> int:
    try:
        return int(float(value or 0))
    except (TypeError, ValueError):
        return 0


def normalize_rows(payload: dict[str, Any], top_n: int) -> list[dict[str, Any]]:
    data = payload.get("data")
    if not isinstance(data, dict):
        raise RuntimeError("接口响应缺少 data 字段。")

    items = data.get("list")
    if not isinstance(items, list) or not items:
        raise RuntimeError("接口响应中没有视频列表。")

    captured_at = datetime.now().astimezone().replace(tzinfo=None)
    rows: list[dict[str, Any]] = []

    for index, item in enumerate(items[:top_n], start=1):
        if not isinstance(item, dict):
            continue
        stat = item.get("stat") if isinstance(item.get("stat"), dict) else {}
        owner = item.get("owner") if isinstance(item.get("owner"), dict) else {}
        cover = normalize_cover_url(item.get("pic"))
        bvid = clean_text(item.get("bvid"))
        aid = to_int(item.get("aid"))
        if bvid:
            link = f"https://www.bilibili.com/video/{bvid}"
        elif aid:
            link = f"https://www.bilibili.com/video/av{aid}"
        else:
            link = ""

        rows.append(
            {
                "captured_at": captured_at,
                "rank": index,
                "cover": safe_excel_text(cover),
                "title": safe_excel_text(item.get("title")),
                "partition": safe_excel_text(
                    item.get("tname")
                    or item.get("typename")
                    or item.get("type_name")
                    or ""
                ),
                "view": to_int(stat.get("view")),
                "like": to_int(stat.get("like")),
                "coin": to_int(stat.get("coin")),
                "favorite": to_int(stat.get("favorite")),
                "share": to_int(stat.get("share")),
                "bvid": safe_excel_text(bvid),
                "owner": safe_excel_text(owner.get("name")),
                "link": link,
            }
        )

    if not rows:
        raise RuntimeError("视频列表中没有任何有效数据。")
    return rows


def apply_header_style(sheet) -> None:
    for cell in sheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = HEADER_ALIGNMENT
        cell.border = THIN_BORDER
    sheet.row_dimensions[1].height = 24
    sheet.freeze_panes = "A2"

    for index, width in enumerate(COLUMN_WIDTHS, start=1):
        sheet.column_dimensions[chr(64 + index)].width = width

    last_column = chr(64 + len(HEADERS))
    sheet.auto_filter.ref = f"A1:{last_column}{max(sheet.max_row, 1)}"


def write_header(sheet) -> None:
    for column, header in enumerate(HEADERS, start=1):
        sheet.cell(1, column, header)
    apply_header_style(sheet)


def append_rows(sheet, rows: list[dict[str, Any]]) -> None:
    first_row_empty = all(
        sheet.cell(1, column).value is None
        for column in range(1, len(HEADERS) + 1)
    )
    excel_row = 1 if first_row_empty else sheet.max_row + 1

    for row in rows:
        rank_change = row.get("rank_change")
        if rank_change is None:
            rank_change_text = "NEW"
        elif rank_change > 0:
            rank_change_text = f"↑{rank_change}"
        elif rank_change < 0:
            rank_change_text = f"↓{abs(rank_change)}"
        else:
            rank_change_text = "—"

        values = [
            row["captured_at"],
            row["rank"],
            rank_change_text,
            row.get("cover", ""),
            row["title"],
            row["partition"],
            row["view"],
            row["like"],
            row["coin"],
            row["favorite"],
            row.get("share", 0),
            row["bvid"],
            row["owner"],
            row["link"],
        ]
        for column, value in enumerate(values, start=1):
            sheet.cell(excel_row, column, value)

        sheet.cell(excel_row, 1).number_format = "yyyy-mm-dd hh:mm:ss"

        for column in NUMERIC_COLUMNS:
            sheet.cell(excel_row, column).number_format = "#,##0"

        for column in range(1, len(HEADERS) + 1):
            cell = sheet.cell(excel_row, column)
            cell.border = THIN_BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=(column == 5))

        rank_change_cell = sheet.cell(excel_row, 3)
        if rank_change is None:
            rank_change_cell.font = Font(color="7A5AF8", bold=True)
        elif rank_change > 0:
            rank_change_cell.font = Font(color="C00000", bold=True)
        elif rank_change < 0:
            rank_change_cell.font = Font(color="008000", bold=True)
        else:
            rank_change_cell.font = Font(color="808080", bold=True)

        # URL 保留为普通文本，避免 openpyxl 在历史记录重复追加时出现关系 ID 冲突。
        excel_row += 1

    last_column = chr(64 + len(HEADERS))
    sheet.auto_filter.ref = f"A1:{last_column}{sheet.max_row}"


def _header_matches(sheet, expected: tuple[str, ...]) -> bool:
    return all(
        sheet.cell(1, column).value == header
        for column, header in enumerate(expected, start=1)
    )


def ensure_history_header(sheet) -> None:
    """确保历史表使用最新表头，并迁移旧版 11/13 列结构。"""
    header_row: int | None = None
    for row_index in range(1, min(sheet.max_row, 10) + 1):
        if sheet.cell(row_index, 1).value == HEADERS[0]:
            header_row = row_index
            break

    if header_row is not None:
        if header_row > 1:
            sheet.delete_rows(1, header_row - 1)

        if _header_matches(sheet, HEADERS):
            apply_header_style(sheet)
            return

        if _header_matches(sheet, OLD_HEADERS_V2):
            # 13 列版本 -> 在“排名”和“封面链接”之间插入“排名变化”。
            sheet.insert_cols(3)
        elif _header_matches(sheet, OLD_HEADERS_V1):
            # 11 列版本 -> 标题前加封面，收藏后加转发，排名后加变化。
            sheet.insert_cols(3)
            sheet.insert_cols(10)
            sheet.insert_cols(3)

        for column, header in enumerate(HEADERS, start=1):
            sheet.cell(1, column, header)
        apply_header_style(sheet)
        return

    first_row_empty = sheet.max_row == 1 and all(
        sheet.cell(1, column).value is None
        for column in range(1, len(HEADERS) + 1)
    )
    if first_row_empty:
        write_header(sheet)
        return

    sheet.insert_rows(1)
    for column, header in enumerate(HEADERS, start=1):
        sheet.cell(1, column, header)
    apply_header_style(sheet)


def write_excel_snapshots(
    snapshots: list[dict[str, Any]],
    output_file: Path,
) -> None:
    if not snapshots:
        return

    output_file.parent.mkdir(parents=True, exist_ok=True)

    try:
        if output_file.exists():
            workbook = load_workbook(output_file)
        else:
            workbook = Workbook()
            workbook.remove(workbook.active)
    except PermissionError as exc:
        raise ExcelWriteError(
            f"无法读取 {output_file}，文件正被 Excel 或其他程序占用。",
            locked=True,
        ) from exc
    except OSError as exc:
        raise ExcelWriteError(f"无法读取 {output_file}：{exc}") from exc

    latest_rows = snapshot_rows(snapshots[-1])
    if "最新排行" in workbook.sheetnames:
        del workbook["最新排行"]
    latest_sheet = workbook.create_sheet("最新排行", 0)
    write_header(latest_sheet)
    append_rows(latest_sheet, latest_rows)

    if "历史记录" not in workbook.sheetnames:
        history_sheet = workbook.create_sheet("历史记录")
    else:
        history_sheet = workbook["历史记录"]

    ensure_history_header(history_sheet)
    for snapshot in snapshots:
        append_rows(history_sheet, snapshot_rows(snapshot))

    workbook.active = 0
    temp_file = output_file.with_name(f".{output_file.stem}.{os.getpid()}.tmp.xlsx")
    try:
        workbook.save(temp_file)
        os.replace(temp_file, output_file)
    except PermissionError as exc:
        raise ExcelWriteError(
            f"无法写入 {output_file}，文件正被 Excel 或其他程序占用。",
            locked=True,
        ) from exc
    except OSError as exc:
        raise ExcelWriteError(
            f"无法写入 {output_file}：{exc}",
            locked=False,
        ) from exc
    finally:
        workbook.close()
        if temp_file.exists():
            temp_file.unlink(missing_ok=True)


def write_excel(rows: list[dict[str, Any]], output_file: Path) -> None:
    """兼容单次调用：将一组行作为快照写入 Excel。"""
    write_excel_snapshots([build_snapshot(rows)], output_file)


def create_dashboard_runtime(settings: Settings) -> DashboardRuntime | None:
    if not settings.dashboard_enabled:
        return None
    return DashboardRuntime(
        html_file=settings.dashboard_html,
        data_file=settings.dashboard_data,
        shortcut_file=settings.dashboard_shortcut,
        refresh_seconds=settings.dashboard_refresh_seconds,
    )


def update_dashboard_safely(
    runtime: DashboardRuntime | None,
    snapshot: dict[str, Any],
    excel_status: str,
    pending_count: int,
    excel_message: str = "",
) -> None:
    if runtime is None:
        return
    try:
        runtime.update(
            snapshot=snapshot,
            excel_status=excel_status,
            pending_count=pending_count,
            excel_message=excel_message,
        )
    except Exception:
        logging.exception("HTML 看板生成失败，但原始数据仍保存在待补写队列中。")


def notify_excel_locked(settings: Settings) -> None:
    global _EXCEL_NOTIFICATION_LAST_SENT

    if not settings.notification_enabled:
        return

    now = time.monotonic()
    if (
        settings.notification_cooldown_seconds > 0
        and now - _EXCEL_NOTIFICATION_LAST_SENT
        < settings.notification_cooldown_seconds
    ):
        return
    _EXCEL_NOTIFICATION_LAST_SENT = now

    prompt = "请关闭 Excel 进行更新。"
    detail = (
        "检测到 Excel 文件正在被占用，本次无法写入。\n\n"
        "请关闭 Excel 进行更新。\n"
        "数据已安全保存在待补写队列中，不会丢失。"
    )
    logging.warning("系统提示：%s", prompt)

    if os.name != "nt":
        notify_send = shutil.which("notify-send")
        if notify_send:
            try:
                subprocess.Popen(
                    [notify_send, "B站热门视频更新提醒", detail],
                    close_fds=True,
                )
            except Exception:
                logging.exception("发送桌面通知失败。")
        return

    powershell = shutil.which("powershell.exe") or "powershell.exe"
    script = (
        "Add-Type -AssemblyName System.Windows.Forms;"
        "$message='检测到 Excel 文件正在被占用，本次无法写入。"
        "`n`n请关闭 Excel 进行更新。"
        "`n数据已安全保存在待补写队列中，不会丢失。';"
        "[void][System.Windows.Forms.MessageBox]::Show("
        "$message,'B站热门视频更新提醒',"
        "[System.Windows.Forms.MessageBoxButtons]::OK,"
        "[System.Windows.Forms.MessageBoxIcon]::Warning)"
    )
    encoded = base64.b64encode(script.encode("utf-16le")).decode("ascii")
    creation_flags = getattr(subprocess, "CREATE_NO_WINDOW", 0)

    try:
        subprocess.Popen(
            [
                powershell,
                "-NoProfile",
                "-NonInteractive",
                "-WindowStyle",
                "Hidden",
                "-EncodedCommand",
                encoded,
            ],
            creationflags=creation_flags,
            close_fds=True,
        )
    except Exception:
        logging.exception("无法弹出 Windows 系统提示。")


def flush_pending_snapshots(
    settings: Settings,
    runtime: DashboardRuntime | None,
) -> int:
    pending = read_pending_snapshots(settings.pending_file)
    if not pending:
        return 0

    try:
        write_excel_snapshots(pending, settings.output_file)
    except Exception as exc:
        update_dashboard_safely(
            runtime,
            pending[-1],
            "failed",
            len(pending),
            str(exc),
        )
        if isinstance(exc, ExcelWriteError) and exc.locked:
            notify_excel_locked(settings)
        raise

    clear_pending_snapshots(settings.pending_file)
    update_dashboard_safely(runtime, pending[-1], "ok", 0)
    logging.info(
        "已将 %d 个待补写快照写入 Excel：%s",
        len(pending),
        settings.output_file,
    )
    return len(pending)


def try_flush_pending(settings: Settings, runtime: DashboardRuntime | None) -> None:
    try:
        flush_pending_snapshots(settings, runtime)
    except PendingQueueError:
        logging.exception("待补写队列损坏，已停止自动清空，请检查 %s", settings.pending_file)
    except Exception:
        logging.warning(
            "Excel 仍有待补写数据，已保留在 %s，下个周期继续尝试。",
            settings.pending_file,
        )


def load_previous_rank_map(settings: Settings) -> dict[str, int]:
    if settings.dashboard_data.exists():
        try:
            payload = json.loads(settings.dashboard_data.read_text(encoding="utf-8"))
            result: dict[str, int] = {}
            for video in payload.get("videos", []):
                bvid = str(video.get("bvid") or "")
                rank = int(video.get("rank") or 0)
                if bvid and rank > 0:
                    result[bvid] = rank
            if result:
                return result
        except Exception:
            logging.warning("无法读取上一轮看板数据，将尝试从待补写队列计算排名变化。")

    try:
        pending = read_pending_snapshots(settings.pending_file)
    except PendingQueueError:
        return {}
    if not pending:
        return {}

    result = {}
    for row in snapshot_rows(pending[-1]):
        bvid = str(row.get("bvid") or "")
        rank = int(row.get("rank") or 0)
        if bvid and rank > 0:
            result[bvid] = rank
    return result


def apply_rank_changes(
    rows: list[dict[str, Any]],
    previous_rank_map: dict[str, int],
) -> list[dict[str, Any]]:
    for row in rows:
        bvid = str(row.get("bvid") or "")
        previous_rank = previous_rank_map.get(bvid)
        row["previous_rank"] = previous_rank
        if previous_rank is None:
            row["rank_change"] = None
        else:
            row["rank_change"] = previous_rank - int(row.get("rank") or 0)
    return rows


def run_once(settings: Settings, runtime: DashboardRuntime | None) -> int:
    logging.info("正在获取 B 站全站热门排行……")
    payload = fetch_ranking(settings)
    rows = normalize_rows(payload, settings.top_n)
    previous_rank_map = load_previous_rank_map(settings)
    rows = apply_rank_changes(rows, previous_rank_map)
    snapshot = build_snapshot(rows)
    pending = append_pending_snapshot(settings.pending_file, snapshot)
    logging.info(
        "抓取成功：%d 条；数据已持久化，待补写队列中有 %d 个快照。",
        len(rows),
        len(pending),
    )
    update_dashboard_safely(
        runtime,
        snapshot,
        "pending",
        len(pending),
        "正在尝试同步 Excel",
    )

    try:
        flush_pending_snapshots(settings, runtime)
    except Exception:
        logging.exception(
            "Excel 写入失败，本次及之前待补写数据不会丢失；"
            "数据已保存在 %s，下个周期会自动重试。",
            settings.pending_file,
        )
    return len(rows)


def main(argv: Sequence[str] | None = None) -> int:
    args = parse_args(argv)
    try:
        settings, config_path = load_settings(args)
    except Exception as exc:
        print(f"配置错误：{exc}", file=sys.stderr)
        return 2

    if args.print_config:
        print(json.dumps(public_settings(settings), ensure_ascii=False, indent=2))
        return 0

    configure_logging(settings.log_level)
    if not config_path.exists():
        logging.warning("未找到配置文件 %s，将使用内置默认值。", config_path)
    logging.info(
        "程序启动：每 %.2f 分钟抓取一次，默认前 %d 条，输出到 %s",
        settings.interval_minutes,
        settings.top_n,
        settings.output_file,
    )
    logging.info("Cookie：%s", "已设置" if settings.cookie else "未设置（公开接口通常不需要）")
    runtime = create_dashboard_runtime(settings)
    if runtime:
        logging.info("HTML 看板文件：%s", settings.dashboard_html)
    else:
        logging.info("HTML 看板已禁用。")

    if args.flush_pending:
        try:
            pending = read_pending_snapshots(settings.pending_file)
            if not pending:
                logging.info("当前没有待补写数据。")
                return 0
            update_dashboard_safely(
                runtime,
                pending[-1],
                "pending",
                len(pending),
                "正在执行手动补写",
            )
            flush_pending_snapshots(settings, runtime)
            return 0
        except KeyboardInterrupt:
            logging.info("用户已停止程序。")
            return 130
        except Exception:
            logging.exception("手动补写失败，数据仍然保留。")
            return 1

    if args.once:
        try:
            run_once(settings, runtime)
            return 0
        except KeyboardInterrupt:
            logging.info("用户已停止程序。")
            return 130
        except Exception:
            logging.exception("本次执行失败。")
            return 1

    server: DashboardHTTPServer | None = None
    refresh_event = threading.Event()
    try:
        if runtime:
            try:
                server, dashboard_url = start_dashboard_server(
                    host=settings.dashboard_host,
                    port=settings.dashboard_port,
                    html_file=settings.dashboard_html,
                    data_file=settings.dashboard_data,
                    refresh_event=refresh_event,
                )
                runtime.write_shortcut(dashboard_url)
                logging.info("实时看板地址：%s", dashboard_url)
                logging.info("也可双击 bilibili_hot_dashboard.url 打开。")
            except Exception:
                logging.exception(
                    "本地看板服务启动失败；HTML 文件仍会更新，可直接打开静态快照。"
                )

        while True:
            started_at = time.monotonic()
            try:
                run_once(settings, runtime)
            except KeyboardInterrupt:
                raise
            except Exception:
                logging.exception("本次抓取失败，将在下个周期继续尝试。")
                try_flush_pending(settings, runtime)

            elapsed = time.monotonic() - started_at
            wait_seconds = max(0.0, settings.interval_minutes * 60 - elapsed)
            next_run = datetime.now().astimezone() + timedelta(seconds=wait_seconds)
            logging.info(
                "下次抓取时间：%s（等待 %.1f 秒）",
                next_run.strftime("%Y-%m-%d %H:%M:%S"),
                wait_seconds,
            )
            try:
                if refresh_event.wait(timeout=wait_seconds):
                    refresh_event.clear()
                    logging.info("收到看板“立即刷新”请求，马上开始新一轮抓取。")
                    continue
            except KeyboardInterrupt:
                raise
    except KeyboardInterrupt:
        logging.info("用户已停止程序。")
        return 130
    finally:
        if server is not None:
            server.shutdown()
            server.server_close()

    return 0

if __name__ == "__main__":
    raise SystemExit(main())